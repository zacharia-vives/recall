"""Put the slides and the three films into one pptx.

Each drawn slide goes in full bleed as a picture. The three films go in as
embedded video with the themed still as their poster frame, so a slide that
has not been clicked yet still reads as a slide. Embedded rather than linked
on purpose: a linked film breaks the moment the deck is copied to another
laptop, which is the classic way to lose a demo.

Speaker notes carry what to say and how long to take, so the deck can be
presented from itself.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Emu

HERE = os.path.dirname(os.path.abspath(__file__))
DL = r"C:\Users\zpsja\Downloads"

W_IN, H_IN = 13.3333, 7.5

# (slide name, video file or None, speaker note)
ORDER = [
    ("v_ad", "Recall-ad-warm-17s.mp4",
     "Say nothing. Click, let it run, let the last frame sit for a beat. "
     "17 seconds. Then straight into the title."),
    ("title", None,
     "Recall. A memory store for people who cannot keep one, with a camera "
     "that does the filing. We are four students at VIVES; the app is live "
     "and you can open it on your own phone right now. 20 seconds."),

    ("sec_intro", None, "Divider. Just read the heading and move on."),
    ("who", None,
     "Two people, and they have opposite needs. Her: seventies, lives "
     "alone, mild memory loss, still reads her own post and does not want "
     "managing. Her daughter: works, cannot ring every day, finds out about "
     "the hospital letter a week late. Name the two-interface decision here, "
     "because everything else follows from it. 45 seconds."),
    ("what", None,
     "What it is, in four parts. Land the second bullet hardest: the card is "
     "the BY-PRODUCT. That is the whole idea. 60 seconds."),
    ("why", None,
     "The strongest slide in the intro. Every reminder app needs somebody to "
     "type, and the person who forgets is exactly the person who will not. "
     "End on the line: a magnifier is a free tool, a magnifier that files is "
     "a product. 45 seconds."),
    ("how", None,
     "Mechanics, briskly. Do not linger, the demo film shows all of this. "
     "40 seconds."),
    ("team", None,
     "Four of us, and say what each person owns. Hand over to Mattiece at "
     "the end of this slide. 30 seconds."),

    ("sec_biz", None, "Divider. Mattiece takes over here."),
    ("model", None,
     "MATTIECE. Two customers on purpose. Be ready for 'why not pick one' "
     "the family model gets users and proves it, the organisation model pays "
     "the salary, and no care organisation buys software with no users. "
     "60 seconds."),
    ("money", "link:https://zacharia-vives.github.io/recall/m/a5d1392bbee7da00d8/money-over-time.html",
     "MATTIECE. Year one and year three. Say the break-even line out loud "
     "before anybody asks: 752 households at the price we can sell today, "
     "537 once multiple keepers ship. Volunteering that is worth more than "
     "defending it. 60 seconds."),
    ("funding", "link:https://zacharia-vives.github.io/recall/m/a5d1392bbee7da00d8/funding-map.html",
     "MATTIECE or ZACHARIA. Calls, not programmes. The one that matters is "
     "imec.istart, which closes 30 September. Finish on the Caring "
     "Technology Principles line, it is the one the Belgian care sector "
     "recognises. 45 seconds."),
    ("campaign1", None,
     "JULIA. The campaign, where older people already are: pharmacies, the "
     "doctor's waiting room, magazines, television. Every advert carries a "
     "QR code. 30 seconds."),
    ("campaign2", None,
     "JULIA. And the second audience: 18 to 50 on Instagram, Facebook, "
     "YouTube and Spotify, because they are the ones who install it for "
     "their parents. 30 seconds."),
    ("campaign3", None,
     "JULIA. What the creative actually looks like. 20 seconds."),
    ("swot", None,
     "JULIA. Her SWOT. Do not read all twenty lines: pick one per box and "
     "move to the next slide, which is where the interesting half is. "
     "45 seconds."),
    ("swot_build", None,
     "JULIA into ZACHARIA. What the build adds. The letter never leaving the "
     "phone is the differentiator no competitor using cloud recognition can "
     "claim. Then own the two weaknesses out loud, including that nobody "
     "over seventy has used it yet. 60 seconds."),

    ("sec_feat", None, "Divider. Zacharia takes over."),
    ("features", None,
     "Do not read this list. Say 'all of this is built and live' and name "
     "four or five. The list exists so the jury can see the scope at a "
     "glance. 30 seconds."),
    ("shots_keeper", None,
     "Her side. Point at the middle screen and say: nothing is filed "
     "silently, and the date is confirmed by hand every time. 30 seconds."),
    ("shots_more", None,
     "Checklists from a photograph, the way there, the voice. 25 seconds."),
    ("shots_helper", None,
     "The family's side on a laptop, and the line under it: no location, no "
     "usage times, no read receipts. Helping is not watching. 35 seconds."),
    ("shots_link", None,
     "Linking and consent. She never sees a login screen: the phone gets an "
     "identity with no password. 30 seconds."),

    ("themes", None,
     "Twenty six looks, six shown. The point is not the colours: it is that "
     "they are generated, with contrast checked by role before the CSS "
     "exists, so a look cannot ship below 4.5:1. Low vision is a named "
     "group, not an afterthought. 30 seconds."),

    ("v_demo", "Recall-demo.mp4",
     "Click and stop talking. 1:22. This is a walkthrough of both sides, not "
     "an advert. If the room is running long, this is the one thing you do "
     "NOT cut."),

    ("sec_gdpr", None, "Divider. This is the section that wins or loses a "
                       "technical jury."),
    ("gdpr_where", None,
     "Where the data is. Say Frankfurt, and say eu-central-1. On her phone "
     "first; the shared copy in Germany; encrypted in transit and at rest; "
     "and the reading and the voice never leave the device. Be precise about "
     "the phone: it is a lock on the screen, not encryption of the cards, "
     "and we say so. 60 seconds."),
    ("gdpr_consent", None,
     "Article 9, because a cardiology letter is health data, so it has to be "
     "explicit consent. Asked out loud, in her language, with the version of "
     "the words recorded. And 7(3): withdrawal is enforced in Postgres by "
     "ten policies, not by the app. 75 seconds. This is the best slide in "
     "the deck, take the time."),
    ("gdpr_less", None,
     "What we chose NOT to collect. This is the strongest thing we can say "
     "and it takes ten seconds: no location, no usage times, no read "
     "receipts. Then the national number being stripped before storage. "
     "45 seconds."),
    ("gdpr_rights", None,
     "The rights, answered by the product rather than by a policy document. "
     "Finish on the honest gaps, unprompted: no key of her own yet, the "
     "library still comes from a CDN, no processor agreement because there "
     "is no company. Volunteering those is what makes the rest credible. "
     "60 seconds."),

    ("sec_notion", None, "Divider."),
    ("notion_what", None,
     "What is on it. The test plan and the process models are the two pages "
     "worth naming. 40 seconds."),
    ("notion_how", None,
     "How we actually used it: Kanban during the session, Daily log at the "
     "end of it, Gantt when somebody asks about Friday. Land the last two: "
     "decisions written down when taken, and it was kept up daily rather "
     "than written the night before. 45 seconds."),

    ("sec_road", None, "Divider."),
    ("roadmap", None,
     "Everything in one place. Lead with itsme for legal documents, because "
     "it is the one with a real reason: proving who actually signed. Then "
     "the social media integrations. Say the WhatsApp limit out loud, it "
     "buys credibility. Then the rest briskly. 75 seconds."),
    ("learning", None,
     "The learning platform, one per side. The argument is the third bullet: "
     "a progress bar tells somebody who forgets that she has forgotten. That "
     "is the sentence that shows we thought about the user rather than the "
     "feature. 45 seconds."),

    ("v_close", "Recall-features.mp4",
     "Click, let it run, do not talk over it. 33 seconds. Then the thank you "
     "slide and questions."),
    ("thanks", None,
     "Thank you. Say the app is live and offline right now, and that there "
     "is an annexe if they want the testing, the security or the numbers."),

    ("sec_annexe", None, "ANNEXE. Only from here on if asked."),
    ("a_testplan", None, "If asked about testing. 407 checks, nine suites, "
                         "five levels, and two things not tested."),
    ("a_defects", None, "If asked what went wrong. The four criticals, and "
                        "the pattern behind five of them."),
    ("a_security", None, "If asked how the data is held. Row level security, "
                         "one gate written once."),
    ("a_arch", None, "If asked about the stack. No framework, no build step, "
                     "local first, and why."),
    ("a_swagger", None,
     "If asked about the API, or if a developer on the jury asks for a spec. "
     "docs/openapi.yaml is a real OpenAPI 3.0.3 document and that is real "
     "Swagger UI rendering it, not a picture of Swagger UI. The one to point "
     "at is the error: a partner call fails the moment she withdraws "
     "consent."),
    ("a_a11y", None, "If asked about accessibility. Zero Axe violations, and "
                     "the honest gap."),
    ("a_costs", None, "If asked about unit economics. Three cents a household "
                      "a month, and why identity changes the shape of that."),
    ("a_compet", None, "If asked who else does this. None of them file."),
    ("a_docs", None, "If asked where anything is written down."),
]


def numbered(name, index, total):
    """A copy of the still with its number on it. Content slides only: a
    number on the opening film or the title reads as a page proof."""
    import sys
    sys.path.insert(0, r"C:\Users\zpsja\AppData\Local\Temp\recall-ad")
    import warm
    from PIL import Image, ImageDraw
    src = os.path.join(HERE, "s_%s.png" % name)
    if name.startswith(("v_", "sec_")) or name in ("title", "thanks",
                                                   "campaign1"):
        return src
    im = Image.open(src).convert("RGB")
    d = ImageDraw.Draw(im)
    d.text((warm.W - 116, warm.H - 54), "%d" % index,
           font=warm.body(26, 500), fill=(214, 138, 96), anchor="rs")
    out = os.path.join(HERE, "n_%s.png" % name)
    im.save(out)
    return out


def build(out="Recall-end-presentation.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(W_IN)
    prs.slide_height = Inches(H_IN)
    blank = prs.slide_layouts[6]

    made = 0
    for name, movie, note in ORDER:
        raw = os.path.join(HERE, "s_%s.png" % name)
        if not os.path.exists(raw):
            print("  MISSING", raw)
            continue
        still = numbered(name, made + 1, len(ORDER))
        slide = prs.slides.add_slide(blank)

        if movie and movie.startswith("link:"):
            # A slide whose picture opens the live page. PowerPoint has no
            # dependable way to run a web page inside a slide, so the click
            # takes the room to the real thing full screen instead, and the
            # capture underneath means the slide still says something when
            # there is no network.
            pic = slide.shapes.add_picture(still, 0, 0, Inches(W_IN),
                                           Inches(H_IN))
            pic.click_action.hyperlink.address = movie[5:]
        elif movie:
            path = os.path.join(DL, movie)
            if not os.path.exists(path):
                print("  MISSING FILM", path)
                slide.shapes.add_picture(still, 0, 0, Inches(W_IN),
                                         Inches(H_IN))
            else:
                # Full bleed, with the drawn still as the poster frame.
                slide.shapes.add_movie(
                    path, 0, 0, Inches(W_IN), Inches(H_IN),
                    poster_frame_image=still, mime_type="video/mp4")
        else:
            slide.shapes.add_picture(still, 0, 0, Inches(W_IN), Inches(H_IN))

        slide.notes_slide.notes_text_frame.text = note
        made += 1

    prs.save(os.path.join(HERE, out))
    print("%s  %d slides" % (out, made))
    return os.path.join(HERE, out)


if __name__ == "__main__":
    build()
